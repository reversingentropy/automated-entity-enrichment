"""The online desk's two halves: what goes to the database, and the page that reads it."""

import json

import pytest


def test_the_deck_is_one_row_per_card_keyed_by_the_card():
    from src.export.deck import deck_rows

    cards = [{"key": "uid:1", "type": "PERSON", "entity": "Koh, Tommy", "pids": [1, 2]},
             {"key": "new:PERSON:x", "type": "PERSON", "entity": "X", "pids": [3]}]
    rows = deck_rows(cards)
    assert [r["key"] for r in rows] == ["uid:1", "new:PERSON:x"]
    assert rows[0]["card"] is cards[0] and rows[0]["type"] == "PERSON" and rows[0]["entity"] == "Koh, Tommy"


def test_the_desk_file_carries_the_index_the_npts_and_the_noted_list():
    from src.export.deck import desk_file

    index = [["1", "Koh, Tommy", "PERSON", "Diplomat", ""]]
    variants = [["许通美", "1", "PERSON"], ["Tommy Koh", "1", "PERSON"]]
    body = json.loads(desk_file(index, variants, [{"entity": "N"}]).decode("utf-8"))
    assert body["index"] == index
    # the record's own words in another order are not an NPT worth shipping
    assert body["npt"] == {"1": ["许通美"]}
    assert body["noted"] == [{"entity": "N"}]


def test_the_online_page_carries_the_client_and_the_keys_but_no_cards():
    from src.export.site import SUPABASE_JS, online_page

    template = ("<script>window.__CARDS__=[];window.__INDEX__=[];window.__NPT__={};window.__NOTED__=[];</script>\n"
                "<script>const ONLINE = window.__ONLINE__ || null;</script>")
    page = online_page(template, "https://p.supabase.co", "anon")
    assert f'<script src="{SUPABASE_JS}"></script>' in page
    assert 'window.__ONLINE__ = {"url": "https://p.supabase.co", "anonKey": "anon"};' in page
    # the library and the keys come before the page's own script
    assert page.index(SUPABASE_JS) < page.index("window.__ONLINE__ =") < page.index("const ONLINE")
    assert page.lstrip().lower().startswith("<!doctype html>") and '<meta charset="utf-8">' in page
    assert "window.__CARDS__=[]" in page


def test_a_test_accounts_decisions_are_left_out(monkeypatch):
    from types import SimpleNamespace
    from src.export import reviews

    users = [SimpleNamespace(id="u1", email="glenn@nlb.gov.sg", user_metadata={"name": "Glenn"}),
             SimpleNamespace(id="t1", email="test@nlb.gov.sg", user_metadata={"test": True})]
    rows = [{"pid": 1, "reviewer": "u1", "decision": {"verdict": "approve"}, "at": "2026-09-18T00:00:00Z"},
            {"pid": 1, "reviewer": "t1", "decision": {"verdict": "reject"}, "at": "2026-09-18T00:00:00Z"}]

    class Admin:
        def list_users(self): return users
    client = SimpleNamespace(auth=SimpleNamespace(admin=Admin()), from_=lambda t: None)
    monkeypatch.setattr("src.shared.supabase_client.get_client", lambda: client)
    monkeypatch.setattr("src.shared.pagination.fetch_all", lambda build, **kw: rows)
    docs = reviews.from_table()
    assert len(docs) == 1
    doc, _ = docs[0]
    assert doc["reviewer"] == "Glenn" and doc["pid"] == 1 and doc["verdict"] == "approve"


def test_an_account_is_made_confirmed_with_its_name_and_a_fresh_password(monkeypatch):
    from types import SimpleNamespace
    from src.export import accounts

    seen = {}

    class Admin:
        def create_user(self, attrs):
            seen.update(attrs)
            return SimpleNamespace(user=SimpleNamespace(id="u9"))
    monkeypatch.setattr(accounts, "get_client", lambda: SimpleNamespace(auth=SimpleNamespace(admin=Admin())))
    made = accounts.create("glenn@nlb.gov.sg", "Glenn Hong")
    assert seen["email_confirm"] is True and seen["user_metadata"] == {"name": "Glenn Hong"}
    assert seen["app_metadata"] == {"approved": True}             # made by key: it does not wait (sql/20)
    assert made["password"] == seen["password"] and len(made["password"]) == 14
    seen.clear()
    test = accounts.create("test@nlb.gov.sg", "Test", test=True)
    # The test mark goes where only the service role can write it.
    assert seen["app_metadata"] == {"approved": True, "test": True} and seen["user_metadata"] == {"name": "Test"}
    assert test["password"] != made["password"]


def test_a_forgotten_password_is_replaced_on_the_right_account(monkeypatch):
    from types import SimpleNamespace

    import pytest

    from src.export import accounts

    updated = {}

    class Admin:
        def list_users(self):
            return [SimpleNamespace(id="u1", email="glenn@nlb.gov.sg"),
                    SimpleNamespace(id="u2", email="minhoon@nlb.gov.sg")]

        def update_user_by_id(self, uid, attrs):
            updated[uid] = attrs["password"]
    monkeypatch.setattr(accounts, "get_client", lambda: SimpleNamespace(auth=SimpleNamespace(admin=Admin())))
    done = accounts.reset_password("MinHoon@nlb.gov.sg ")
    assert updated == {"u2": done["password"]} and len(done["password"]) == 14
    with pytest.raises(ValueError):
        accounts.reset_password("nobody@nlb.gov.sg")


def test_a_chosen_password_is_asked_twice_and_never_printed(monkeypatch, capsys):
    from types import SimpleNamespace

    import src.export.__main__ as cli
    from src.export import accounts

    typed = iter(["short", "long enough 1", "long enough 2", "long enough 1", "long enough 1"])
    assert accounts.ask_password(lambda prompt: next(typed)) == "long enough 1"
    out = capsys.readouterr().out
    assert "Too short" in out and "did not match" in out

    seen = {}

    class Admin:
        def create_user(self, attrs):
            seen.update(attrs)
            return SimpleNamespace(user=SimpleNamespace(id="u9"))
    monkeypatch.setattr(accounts, "get_client", lambda: SimpleNamespace(auth=SimpleNamespace(admin=Admin())))
    monkeypatch.setattr(accounts.getpass, "getpass", lambda prompt: "my own choice")
    assert cli.main(["--account", "someone@nlb.gov.sg", "Their Name", "--password"]) == 0
    assert seen["password"] == "my own choice" and seen["user_metadata"] == {"name": "Their Name"}
    out = capsys.readouterr().out
    assert "Created someone@nlb.gov.sg (Their Name)" in out and "my own choice" not in out


def test_cards_are_dealt_latest_news_first():
    from src.export.cards import deal_order

    def card(entity, *sources, subject=False):
        return {"entity": entity, "subject": subject, "isFlag": False,
                "sources": [{"date": d, "articleId": a} for d, a in sources]}

    old = card("Old news", ("2026-09-19", 50))
    again = card("Reported again today", ("2026-09-19", 40), ("2026-09-25", 90))
    today_subject = card("Today's subject", ("2026-09-25", 95), subject=True)
    today_mention = card("Today's mention", ("2026-09-25", 95))
    undated = card("No date", ("", 99))
    dealt = sorted([old, undated, today_mention, again, today_subject], key=deal_order)
    assert [c["entity"] for c in dealt] == [
        "Today's subject", "Today's mention", "Reported again today", "Old news", "No date"]


def test_the_weekly_file_is_the_template_agreed_with_the_reviewers(tmp_path):
    from openpyxl import load_workbook

    from src.export.weekly import HEADERS, rows_for, write

    link = "https://www.channelnewsasia.com/singapore/kenneth-jeyaretnam-reform-party-leader-6263591"
    cards = {
        "1": {"key": "uid:18518926", "entity": "Kenneth Jeyaretnam", "type": "PERSON", "pids": [1, 2],
              "sources": [{"link": link}]},
        "3": {"key": "new:PERSON:Amos Yee", "entity": "Amos Yee", "type": "PERSON", "pids": [3],
              "newsFields": [{"key": "Birth Year (yyyy)", "value": "1998"}], "sources": [{"link": "https://example.com"}]},
        "4": {"key": "uid:9", "entity": "Someone", "type": "PERSON", "pids": [4], "sources": [{"link": "x"}]},
    }
    decided = {"verdict": "approve", "chose": "18518926", "record": "Jeyaretnam, Kenneth", "reviewer": "Yaw Huah",
               "rows": [{"field": "Death Year (yyyy)", "current": "", "proposed": "2026", "from": link}]}
    decisions = [
        {**decided, "pid": 1, "at": "2026-09-28T02:00:00Z"},
        # The same card settled a second article's proposal: one row, not two.
        {**decided, "pid": 2, "at": "2026-09-28T02:00:00Z"},
        {"pid": 3, "verdict": "approve", "chose": "new-entity", "reviewer": "Yaw Huah", "at": "2026-09-28T03:00:00Z",
         "done": "2026-09-29T02:00:00Z"},
        # Rejected, or kept for review: nothing to apply.
        {"pid": 4, "verdict": "reject", "chose": None, "reviewer": "Glenn", "at": "2026-09-28T03:00:00Z"},
    ]
    existing, new = rows_for(decisions, cards, {"18518926": "_People"}, {"PERSON": "_People"})
    assert existing == [["18518926", "Jeyaretnam, Kenneth", "_People", "Death Year (yyyy)", "", "2026", link,
                         "Yaw Huah", "28/9/2026", ""]]
    # Marked done on the desk: the Status column says so.
    assert new == [["", "Amos Yee", "_People", "Birth Year (yyyy)", "", "1998", "https://example.com",
                    "Yaw Huah", "28/9/2026", "Done"]]

    wb = load_workbook(write(existing, new, tmp_path / "week.xlsx"))
    assert wb.sheetnames == ["Existing entities", "New entities"]
    assert [c.value for c in wb["Existing entities"][1]] == HEADERS and HEADERS[-1] == "Status"
    assert wb["New entities"]["B2"].value == "Amos Yee"


def test_a_record_to_create_goes_on_the_sheet_as_the_reviewer_left_it():
    from src.export.weekly import rows_for

    cards = {"5": {"key": "new:ORGANISATION:Kampung Senang", "entity": "Kampung Senang", "type": "ORGANISATION", "pids": [5],
                   "sources": [{"link": "https://example.com"}],
                   "newsFields": [{"key": "Affiliations(groupName)", "value": "People's Association | Ministry of Health"},
                                  {"key": "Year Started (yyyy)", "value": "1999"}]}}
    # The reviewer made it a programme and left out one affiliation and the year with ×.
    decisions = [{"pid": 5, "verdict": "approve", "chose": "new-entity", "reviewer": "R", "at": "2026-10-24T03:00:00Z",
                  "ctype": "PROGRAMME", "leftOut": [{"field": "Affiliations(groupName)", "value": "Ministry of Health"},
                                                   {"field": "Year Started (yyyy)", "value": "1999"}]}]
    _, new = rows_for(decisions, cards, {}, {"ORGANISATION": "_Organisations", "PROGRAMME": "_Programmes"})
    assert [row[2:6] for row in new] == [["_Programmes", "Affiliations(groupName)", "", "People's Association"]]

    # Edited on the desk: the sheet carries the fields as the reviewer left them.
    decisions[0]["fields"] = [{"key": "Affiliations(groupName)", "value": "People's Association (PA)"}]
    _, new = rows_for(decisions, cards, {}, {"PROGRAMME": "_Programmes"})
    assert [row[3:6] for row in new] == [["Affiliations(groupName)", "", "People's Association (PA)"]]


def _desk(n_waiting_people=2):
    cards = [{"key": f"uid:{i}", "type": "PERSON", "entity": f"Person {i}",
              "sources": [{"link": "https://www.straitstimes.com/x"}]} for i in range(n_waiting_people)]
    cards += [{"key": "uid:z", "type": "ORGANISATION", "entity": "Zaobao Org",
               "sources": [{"link": "https://www.zaobao.com.sg/x"}]}]
    return cards


def test_the_daily_email_is_silent_when_there_is_nothing_to_do():
    from src.export.notify import messages

    glenn = {"id": "g", "email": "glenn@nlb.gov.sg", "name": "Glenn", "langs": ["en"], "types": None}
    cards = _desk(0)   # only a Chinese card, which Glenn does not read
    assert messages([glenn], cards, [], [], {"new_cards": 0, "dump": None, "failed": []}) == []


def test_the_daily_email_carries_each_reviewers_own_list_and_only_team_totals():
    from src.export.notify import messages

    glenn = {"id": "g", "email": "glenn@nlb.gov.sg", "name": "Glenn", "langs": ["en"], "types": ["people"]}
    min_hoon = {"id": "m", "email": "mh@nlb.gov.sg", "name": "Min Hoon", "langs": ["en", "zh"], "types": None}
    cards = _desk(3)
    reviews = [
        {"pid": 1, "card_key": "uid:0", "reviewer": "g",
         "decision": {"verdict": "approve", "rows": [{"field": "Awards"}, {"field": "Description"}]}},
        {"pid": 2, "card_key": "uid:1", "reviewer": "m", "decision": {"verdict": "approve", "done": "2026-09-28"}},
    ]
    holds = [{"card_key": "uid:2", "reviewer": "g"}]
    mails = {m["to"]: m for m in messages([glenn, min_hoon], cards, reviews, holds,
                                          {"new_cards": 0, "dump": None, "failed": []})}
    g = mails["glenn@nlb.gov.sg"]
    assert "On your desk: Person 2" in g["body"]
    assert "Person 0 (2 changes)" in g["body"]                  # his approval, not yet done
    assert "Waiting" not in g["body"]                           # nothing left in his part: people only, English
    # Min Hoon's approval is done and nothing is on her desk, but a Chinese card waits for her.
    m = mails["mh@nlb.gov.sg"]
    assert "Waiting in your part of the pile: 1 (Organisations 1)" in m["body"]
    assert "To make in TTE" not in m["body"] and "Glenn" not in m["body"]
    assert "The team so far: 4 cards, 2 decided, 0 kept for review." in m["body"]


def test_news_alone_is_worth_an_email():
    from src.export.notify import messages

    glenn = {"id": "g", "email": "glenn@nlb.gov.sg", "name": "Glenn", "langs": ["en"], "types": None}
    mails = messages([glenn], _desk(0), [], [], {"new_cards": 0, "dump": "TTE-DELTA_20261101.zip", "failed": ["extraction"]})
    assert len(mails) == 1
    assert "TTE dump TTE-DELTA_20261101.zip loaded; the extraction step failed" in mails[0]["body"]


def test_a_site_build_refuses_the_url_and_the_key_swapped(tmp_path, monkeypatch):
    from src.export.site import build_site

    template = tmp_path / "t.html"
    template.write_text("<html><head></head><body></body></html>", encoding="utf-8")
    monkeypatch.setenv("SUPABASE_URL", "eyJhbGciOi.payload.signature")
    monkeypatch.setenv("SUPABASE_ANON_KEY", "https://p.supabase.co")
    with pytest.raises(RuntimeError, match="swapped"):
        build_site(template, tmp_path / "out")
    monkeypatch.setenv("SUPABASE_URL", " https://p.supabase.co \n")
    monkeypatch.setenv("SUPABASE_ANON_KEY", "eyJhbGciOi.payload.signature")
    assert build_site(template, tmp_path / "out").exists()


def test_unsent_mail_is_not_printed_into_a_public_run_log(monkeypatch, capsys):
    from src.export.notify import send

    mail = [{"to": "someone@nlb.gov.sg", "subject": "TTE desk", "body": "On your desk: a record."}]
    for k in ("MAIL_SERVER", "MAIL_FROM", "MAIL_USERNAME", "MAIL_PASSWORD"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("GITHUB_ACTIONS", "true")
    assert send(mail) == 0
    out = capsys.readouterr().out
    assert "someone@nlb.gov.sg" not in out and "On your desk" not in out and "1 message" in out
    assert send(mail, dry_run=True) == 0 and "someone@nlb.gov.sg" in capsys.readouterr().out   # by hand, shown
    monkeypatch.delenv("GITHUB_ACTIONS")
    send(mail)
    assert "someone@nlb.gov.sg" in capsys.readouterr().out


def test_an_account_made_on_the_page_waits_until_approved_and_gets_no_mail():
    from datetime import datetime, timezone
    from types import SimpleNamespace

    from src.export.accounts import is_waiting

    since = datetime(2026, 10, 23, tzinfo=timezone.utc)
    user = lambda made, **meta: SimpleNamespace(created_at=made, app_metadata=meta)
    assert is_waiting(user("2026-10-24T01:00:00Z"), since)
    assert not is_waiting(user("2026-10-24T01:00:00Z", approved=True), since)
    assert not is_waiting(user("2026-10-24T01:00:00Z", test=True), since)
    assert not is_waiting(user("2026-10-01T01:00:00Z"), since)                 # made before sql/20
    assert not is_waiting(user("2026-10-24T01:00:00Z"), None)                  # sql/20 not run yet
