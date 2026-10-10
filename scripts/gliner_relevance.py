"""
Does a small, free classifier agree with the LLMs on relevance?

GLiNER2.5-multi-Decide (Fastino; Apache 2.0; 287M parameters; GPU or CPU) is
asked the relevance question and compared with:

  live     Gemini's verdicts on the live feed          data/gliner/live-articles.json
  archive  the internal model's verdicts, 2015-2026    data/backfill/answers/relevance.csv
  labels   69 human labels, all rows where Gemini and the internal model
           disagreed, English only                     data/backfill/labels/relevance-labels.csv
  sheet    the evaluation's 736 August articles         data/eval/1-relevance.xlsx

Against a model's verdict, precision, recall and F1 measure agreement, not
accuracy: the reference model can be wrong too. Only `labels`, and `sheet`
once the team has filled it in, are scored against people.

Answers are appended to data/gliner/<set>-answers.csv as they arrive, so a
stopped run carries on where it left off, and `report` rescores them without
loading the model. Each report is also written to data/gliner/report-<set>.txt.

  python scripts/gliner_relevance.py live
  python scripts/gliner_relevance.py archive                 # all 295,269
  python scripts/gliner_relevance.py archive --sample 20000  # a fixed random sample
  python scripts/gliner_relevance.py labels
  python scripts/gliner_relevance.py sheet
  python scripts/gliner_relevance.py report live             # live | archive | labels | sheet

The question can be put three ways (--task): `described`, the default, is two
labels carrying the relevance prompt's definition; `events` is one short label
per kind of record change plus "other news"; `short` is two bare labels. Each
keeps its own answers file, so they can be compared on the same articles:

  python scripts/gliner_relevance.py live --task events
  python scripts/gliner_relevance.py report live --task events

Not part of the pipeline: it imports nothing from src/, and its packages
(torch, gliner2 and the rest) are installed separately, not in the project.
"""

import argparse
import bisect
import csv
import html
import json
import random
import sys
import time
from pathlib import Path

csv.field_size_limit(sys.maxsize)
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "gliner"
LIVE = OUT / "live-articles.json"
ARCHIVE = ROOT / "data" / "backfill" / "answers" / "relevance.csv"
LABELS = ROOT / "data" / "backfill" / "labels" / "relevance-labels.csv"
SHEET = ROOT / "data" / "eval" / "1-relevance.xlsx"
MODEL = "fastino/GLiNER2.5-multi-Decide"

# The relevance prompt's definition (prompts/relevance.txt), as two described labels.
DESCRIBED = {"relevance": {"labels": {
    "relevant": (
        "Reports a completed or immediate change to a named Singapore person, organisation, building, place, "
        "event, programme or law: someone died, was appointed, retired or received a named award; an "
        "organisation was created, merged, closed or renamed; a building opened, closed or got heritage status; "
        "a programme was launched or ended; a law was enacted, amended or repealed."),
    "not relevant": (
        "Commentary or opinion, people merely quoted or reacting, routine updates, court cases, foreign "
        "entities, or plans that have not yet taken effect."),
}}}
# The same question as the routing tasks this kind of model is trained on: one
# short label per kind of record change, and everything else as "other news".
EVENTS = {"news": {"labels": {
    "appointment": "a named person was appointed, elected or promoted to a post",
    "departure": "a named person retired, resigned or stepped down from a post",
    "death": "a named person died",
    "award": "a named person or organisation received a named award, honour or scholarship",
    "organisation change": "an organisation was founded, merged, closed, renamed or relocated",
    "building or place change": "a building or place opened, closed, was demolished, renamed or given heritage status",
    "programme change": "a named programme or scheme was launched or ended",
    "law change": "a law or bill was passed, amended or repealed",
    "event change": "a named event was cancelled, postponed, renamed or relocated",
    "other news": ("anything else: commentary, opinion, analysis, crime and courts, business results, sport, "
                   "lifestyle, people quoted or reacting, plans not yet in effect"),
}}}
SHORT = {"relevance": ["reports a change to a named person, organisation, place or law", "other news"]}
# name -> (task key, schema, which answer counts as relevant)
TASKS = {
    "described": ("relevance", DESCRIBED, lambda label: label == "relevant"),
    "events": ("news", EVENTS, lambda label: label != "other news"),
    "short": ("relevance", SHORT, lambda label: label == SHORT["relevance"][0]),
}
REFERENCE = {"live": "Gemini", "archive": "the internal model", "labels": "people", "sheet": "people"}


def lang_of(url, source=""):
    return "zh" if source == "zb" or "zaobao" in (url or "") else "en"


def text_of(title, description):
    # Feed text arrives HTML-escaped ("Singapore&#039;s"); the classifier reads text.
    parts = [html.unescape(str(x)).strip() for x in (title, description) if x]
    return ". ".join(p for p in parts if p)


def truth(value):
    """TRUE / True / true -> True, FALSE -> False, blank -> None."""
    if value is None or isinstance(value, bool):
        return value
    s = str(value).strip().lower()
    return True if s in ("true", "1", "yes") else False if s in ("false", "0", "no") else None


# ---- the four sets: each a list of {key, lang, text, ref} ----------------------

def live_items():
    if not LIVE.exists():
        sys.exit(f"{LIVE} is missing: it is exported from the database, not built here.")
    return [{"key": str(r["id"]), "lang": lang_of(r.get("url")),
             "text": text_of(r.get("title"), r.get("description")), "ref": truth(r.get("relevant"))}
            for r in json.loads(LIVE.read_text(encoding="utf-8"))]


def archive_items():
    with open(ARCHIVE, encoding="utf-8", newline="") as f:
        return [{"key": r["url"], "lang": lang_of(r["url"], r["source"]),
                 "text": text_of(r["title"], r["description"]), "ref": truth(r["relevant"])}
                for r in csv.DictReader(f) if truth(r["relevant"]) is not None]


def label_rows():
    with open(LABELS, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def labels_items():
    rows = label_rows()
    want = {r["key"] for r in rows}
    # The labels file has headlines only; the archive has the descriptions the models saw.
    description = {}
    with open(ARCHIVE, encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            if r["url"] in want:
                description[r["url"]] = r["description"]
    return [{"key": r["key"], "lang": lang_of(r["key"], r["source"]),
             "text": text_of(r["title"], description.get(r["key"])), "ref": truth(r["human"])}
            for r in rows]


def sheet_items():
    from openpyxl import load_workbook
    rows = load_workbook(SHEET, read_only=True)["rows"].iter_rows(values_only=True)
    head = [str(x).strip().lower() if x else "" for x in next(rows)]
    items = []
    for r in rows:
        rec = dict(zip(head, r))
        if rec.get("url"):
            items.append({"key": rec["url"], "lang": lang_of(rec["url"], rec.get("source") or ""),
                          "text": text_of(rec.get("title"), rec.get("description")),
                          "ref": truth(rec.get("relevant"))})
    return items


LOADERS = {"live": live_items, "archive": archive_items, "labels": labels_items, "sheet": sheet_items}


def answers_path(name, task):
    return OUT / (f"{name}-answers.csv" if task == "described" else f"{name}-{task}-answers.csv")


# ---- running the model ---------------------------------------------------------

def load_model(half):
    import torch
    from gliner2 import AutoExtractor

    model = AutoExtractor.from_pretrained(MODEL)
    device = "cuda" if torch.cuda.is_available() else "cpu"
    model = model.to(device)
    if half and device == "cuda":
        model = model.to(torch.bfloat16)
    model.eval()
    print(f"model on {device}" + (" in bfloat16" if half and device == "cuda" else ""), flush=True)
    if device == "cpu":
        print("  no GPU found: expect about 7 articles a second", flush=True)
    return model


def parse(out, key):
    """(label, confidence) from one result, whichever shape the library returns."""
    r = out.get(key, out) if isinstance(out, dict) else out
    if isinstance(r, list):
        r = r[0] if r else {}
    if isinstance(r, dict):
        return r.get("label"), float(r.get("confidence", 1.0))
    return str(r), 1.0


def run(name, items, batch_size, half, task):
    import torch

    key, schema, is_relevant = TASKS[task]
    path = answers_path(name, task)
    done = set()
    if path.exists():
        with open(path, encoding="utf-8", newline="") as f:
            done = {r["key"] for r in csv.DictReader(f)}
    todo = [it for it in items if it["key"] not in done and it["text"]]
    print(f"{name} ({task}): {len(items):,} articles, {len(done):,} already answered, {len(todo):,} to go", flush=True)
    if not todo:
        return
    model = load_model(half)
    OUT.mkdir(parents=True, exist_ok=True)
    fresh = not path.exists()
    started = time.time()
    chunk_size = batch_size * 32
    with open(path, "a", encoding="utf-8", newline="") as f, torch.inference_mode():
        out = csv.writer(f)
        if fresh:
            out.writerow(["key", "lang", "label", "confidence", "p_relevant", "relevant"])
        for i in range(0, len(todo), chunk_size):
            chunk = todo[i:i + chunk_size]
            results = model.batch_classify_text([it["text"] for it in chunk], schema,
                                                batch_size=batch_size, include_confidence=True)
            for it, result in zip(chunk, results):
                label, confidence = parse(result, key)
                relevant = is_relevant(label)
                # A ranking score: the confidence in a relevant answer, or one minus the confidence in an irrelevant one.
                p = confidence if relevant else 1 - confidence
                out.writerow([it["key"], it["lang"], label, f"{confidence:.4f}", f"{p:.4f}", int(relevant)])
            f.flush()
            n = i + len(chunk)
            rate = n / max(time.time() - started, 1e-9)
            print(f"  {len(done) + n:,} / {len(items):,}   {rate:,.0f} a second   "
                  f"about {(len(todo) - n) / rate / 60:.0f} min left", flush=True)


# ---- scoring -------------------------------------------------------------------

def counts(pairs):
    tp = sum(p and r for p, r in pairs)
    fp = sum(p and not r for p, r in pairs)
    fn = sum(r and not p for p, r in pairs)
    tn = len(pairs) - tp - fp - fn
    return tp, fp, fn, tn


def score_line(name, pairs):
    tp, fp, fn, tn = counts(pairs)
    n = len(pairs)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    agree = (tp + tn) / n if n else 0.0
    chance = ((tp + fp) * (tp + fn) + (fn + tn) * (fp + tn)) / (n * n) if n else 0.0
    kappa = (agree - chance) / (1 - chance) if chance < 1 else 0.0
    return (f"{name:<22}{n:>8,}{tp + fn:>8,}{tp + fp:>8,}{tp:>7,}{fp:>7,}{fn:>7,}"
            f"{precision:>10.1%}{recall:>8.1%}{f1:>7.2f}{agree:>8.1%}{kappa:>7.2f}")


HEADER = (f"{'':<22}{'rows':>8}{'ref yes':>8}{'clf yes':>8}{'TP':>7}{'FP':>7}{'FN':>7}"
          f"{'precision':>10}{'recall':>8}{'F1':>7}{'agree':>8}{'kappa':>7}")


def auc(scored):
    pos = sorted(p for p, r in scored if r)
    neg = sorted(p for p, r in scored if not r)
    if not pos or not neg:
        return None
    wins = sum(bisect.bisect_left(neg, x) + 0.5 * (bisect.bisect_right(neg, x) - bisect.bisect_left(neg, x))
               for x in pos)
    return wins / (len(pos) * len(neg))


def report(name, examples=5, task="described"):
    path = answers_path(name, task)
    if not path.exists():
        sys.exit(f"No answers for {name} ({task}) yet: run `python scripts/gliner_relevance.py {name} --task {task}` first.")
    is_relevant = TASKS[task][2]
    with open(path, encoding="utf-8", newline="") as f:
        answers = {r["key"]: r for r in csv.DictReader(f)}
    for a in answers.values():
        a["yes"] = a["relevant"] == "1" if a.get("relevant") not in (None, "") else is_relevant(a["label"])
    items = {it["key"]: it for it in LOADERS[name]()}
    rows = [(items[k], a) for k, a in answers.items() if k in items]
    lines = []
    say = lines.append
    reference = REFERENCE[name]
    scored = [(it, a) for it, a in rows if it["ref"] is not None]
    if not scored:
        called = sum(a["yes"] for _, a in rows)
        say(f"{name}: {len(rows):,} answers kept; nothing labelled yet, so nothing to score.")
        say(f"  The classifier calls {called:,} relevant ({called / max(len(rows), 1):.0%}).")
    else:
        say(f"{name}: GLiNER2.5-multi-Decide, asked as `{task}`, against {reference}"
            + ("" if reference == "people" else " (agreement with a model, not accuracy)"))
        say("  precision: of the articles the classifier calls relevant, the share the reference also does")
        say("  recall: of the articles the reference calls relevant, the share the classifier also catches")
        say("  kappa: agreement beyond what chance would give; 1 is perfect, 0 is chance")
        say("")
        say(HEADER)
        for lang, title in (("all", "all"), ("en", "English"), ("zh", "Chinese")):
            sub = [(a["yes"], it["ref"]) for it, a in scored if lang == "all" or it["lang"] == lang]
            if sub:
                say(score_line(f"classifier, {title}", sub))
        if name == "labels":
            # The two LLMs on the same rows, against the same people.
            by_key = {r["key"]: r for r in label_rows()}
            for model in ("gemini", "internal"):
                pairs = [(truth(by_key[it["key"]][model]), it["ref"]) for it, _ in scored]
                say(score_line(f"{model}, all", pairs))

        # Ranking: how many articles must pass to catch a given share of the reference's relevant ones.
        ranked = sorted(((float(a["p_relevant"]), it["ref"]) for it, a in scored), key=lambda x: -x[0])
        positives = sum(r for _, r in ranked)
        area = auc(ranked)
        say("")
        say(f"Ranking by the classifier's confidence (AUC {area:.2f}):" if area is not None else "Ranking:")
        for target in (0.80, 0.90, 0.95, 0.99):
            caught = 0
            for k, (p, r) in enumerate(ranked, 1):
                caught += r
                if positives and caught >= target * positives:
                    say(f"  to catch {target:.0%} of {reference}'s relevant articles, pass the top {k:,} "
                        f"({k / len(ranked):.0%} of all), scores above {p:.2f}")
                    break

        if task == "events":
            from collections import Counter
            chosen = Counter(a["label"] for _, a in scored)
            say("")
            say("Labels chosen: " + ", ".join(f"{label} {n:,}" for label, n in chosen.most_common()))

        if examples:
            text = lambda it, a: f"  {it['lang']} {float(a['p_relevant']):.2f}  [{a['label']}]  {it['text'][:100]}"
            missed = sorted(((it, a) for it, a in scored if it["ref"] and not a["yes"]),
                            key=lambda x: float(x[1]["p_relevant"]))
            extra = sorted(((it, a) for it, a in scored if not it["ref"] and a["yes"]),
                           key=lambda x: -float(x[1]["p_relevant"]))
            say("")
            say(f"Relevant to {reference}, not to the classifier (least confident first):")
            lines.extend(text(it, a) for it, a in missed[:examples])
            say(f"Relevant to the classifier, not to {reference} (most confident first):")
            lines.extend(text(it, a) for it, a in extra[:examples])

    report_text = "\n".join(lines)
    print(report_text)
    suffix = "" if task == "described" else f"-{task}"
    (OUT / f"report-{name}{suffix}.txt").write_text(report_text + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("set", choices=["live", "archive", "labels", "sheet", "report"])
    parser.add_argument("which", nargs="?", choices=list(LOADERS), help="with report: which set to score")
    parser.add_argument("--sample", type=int, help="a fixed random sample of this many articles")
    parser.add_argument("--batch-size", type=int, default=64)
    parser.add_argument("--half", action="store_true", help="bfloat16 on the GPU: faster, very slightly different")
    parser.add_argument("--examples", type=int, default=5, help="disagreements to list in the report")
    parser.add_argument("--task", choices=list(TASKS), default="described", help="how the question is put")
    args = parser.parse_args()

    if args.set == "report":
        if not args.which:
            parser.error("report needs a set: live, archive, labels or sheet")
        report(args.which, args.examples, args.task)
        return
    items = LOADERS[args.set]()
    if args.sample and args.sample < len(items):
        # Seeded, so a rerun draws the same sample and resumes it.
        items = random.Random(7).sample(items, args.sample)
    run(args.set, items, args.batch_size, args.half, args.task)
    report(args.set, args.examples, args.task)


if __name__ == "__main__":
    main()
