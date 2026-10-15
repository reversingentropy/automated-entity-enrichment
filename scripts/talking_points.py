"""
The talk's speaker notes, corrected against the paper, and a talking-points document with a running clock.

    uv run python scripts/talking_points.py

Reads the deck Ashwin edits (data/dist/designing-for-the-reviewer-v2.pptx) and never writes it. Writes a copy with
only the notes changed (designing-for-the-reviewer-v2-talking-points.pptx) and data/dist/talking-points.docx: every
slide's notes in order, the talk's slides with a running clock at 130 words a minute, the backup slides after.
Slides whose notes are not listed in NOTES keep Ashwin's notes word for word.
"""

import sys
from pathlib import Path

import docx
from docx.shared import Pt
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "data" / "dist" / "designing-for-the-reviewer-v2.pptx"
OUT_DECK = ROOT / "data" / "dist" / "designing-for-the-reviewer-v2-talking-points.pptx"
OUT_DOC = ROOT / "data" / "dist" / "talking-points.docx"
TALK_SLIDES = 19                 # slide 19 is the thank-you; the rest are backup
WPM = 130
PAUSES = {10: 10}                # seconds of silence written into a slide's notes ("Take ten seconds")

TITLES = {1: "Title", 2: "The National Library Board", 3: "The problem", 4: "What is an authority record?", 5: "The labour",
          6: "How it works", 7: "One thread, four checks", 8: "The AI reads, a person decides", 9: "Papers, Please or Tinder?",
          10: "You are the reviewer. Approve?", 11: "The reveal: a career replaced by a compliment",
          12: "A guessed translation can look like a perfect match", 13: "Does anything disagree with this change?",
          14: "What was left out", 15: "Four checks, and what our screen does", 16: "What if we remove the human in the loop?",
          17: "What we do not know, and what you could do", 18: "The four questions, again", 19: "Thank you",
          20: "Backup: without the person, in full", 21: "Backup: without the person, the short version",
          22: "Backup: the numbers", 23: "Backup: the evaluation plan", 24: "Backup: cost, law and retrieval",
          25: "Backup: the schema belongs in code", 26: "Backup: could a small model do the first filter?",
          27: "Backup: more cards from the live deck", 28: "Backup: an earlier deletion, 636 to 405",
          29: "Backup: a missing step, found by timing a reviewer"}

# The notes that change, and why (the rest stay as Ashwin wrote them).
NOTES = {
    # relevance reads the title and the summary, not the title alone (paper, Section 2)
    6: """Here is how it works. Six stages. The first five run on their own, on a schedule.

Ingest collects titles and summaries from three news feeds every two hours. Relevance reads only the title and summary: does this change a record? Extraction pulls out the people, places and facts, with one quoted sentence. Retrieval finds candidate records by character similarity, and uses no AI. Resolution asks: is this the same entity, and what does the record lack?

The sixth stage is a person. Run over our archive back to 2015 and the live feed, it has read almost three hundred thousand articles, judged about twenty-eight thousand relevant, and proposed 9,844 changes to existing records. None of those has been through a formal evaluation yet.""",
    # five cataloguers chose the redesign in October; a little shorter to fit the 15 minutes
    9: """So we changed the screen. If you have played Papers, Please, you know the feeling: you are a border inspector with two documents, you check them against the rules, and you stamp approved or denied. Our earlier desk was inspired by that, and by Tinder: the article on one side, the record on the other, one question, a stamp.

Two cataloguers used it, found problems, and helped redesign it. Then a first-time user, someone outside the library, looked at it and said: it feels like Tinder.

That is the opposite of what we want. Papers, Please makes you slow down and look. Tinder makes you swipe. The same screen can be either, and the design decides which. So the redesign has one card and one question: the news on top, then what the record says now, then what it will say after you accept. In October, five cataloguers tried both versions side by side, and chose the redesign.""",
    # written for this slide, not the old mirrored chart; with the baseline the paper now gives (43 of 48)
    16: """Someone will ask: why keep the person at all? We tried it, on stored data. A second model, jev from TypeSafe, made every decision; it writes no text, so the stored sentences were reused.

On forty-eight cards a reviewer had decided, it reached the reviewer's outcome on sixteen. It decides in milliseconds; the reviewer took a median of thirty-three seconds a card. Faster, but not as good.

But had the stored proposals simply been accepted unreviewed, they would have matched the reviewer on forty-three of the forty-eight. Those proposals came mostly from the archive's internal model. So most of the gap is swapping the model, not removing the person. The five it misses are judgement calls, like the three cards the reviewer thought not worth a record. Taking the person out does not remove those decisions. It moves them into cut-offs that nobody reviews.

A pilot: one reviewer, who is not a cataloguer, on cards chosen partly where the models disagreed.""",
    # slide 16 is one pilot, so "untested" becomes "not settled"; the agentic point moves here from slide 16
    17: """What do we know? Less than I would like. Can a person decide each proposal? Yes. Is each proposal correct? We have not measured that. Does it save time? Not yet. In the pilot, records built from many articles took three quarters of the reviewer's time, so description notes are now capped. Whether a second model could replace the person is not settled, and we have not tried agentic workflows; their costs are unknown.

What would it take for you? One person built this, relying on a team for feedback, on free tiers. Paid access is estimated at under fifteen Singapore dollars a month, about a hundred and ninety-five rand. The hard part was never the model. It was the screen and the steps around it.

And before you start, check your copyright position. Ours is Singapore's computational data analysis provision, and the article is read, not stored. Nothing reaches our file without a named person's decision.""",
    # the backup slides, in their present order
    19: """So: Papers, Please, not Tinder. A human in the loop is only real if the human can see enough to disagree. Build the screen so that a person can look, and so that they can see what was left out.

Thank you. I would be glad to take your questions. The slides after this one are backup, if they help.""",
    # the stored answers are nine in ten the archive's internal model, not Gemini's; "2 not worth" was 1 + 1 kept;
    # the baseline added; the agentic reminder moved to slide 17
    20: """Only if asked. This is the comparison in full. Left is how it runs now, right is the same pipeline without the person: jev from TypeSafe makes the decisions, and the stored sentences are reused.

Relevance, counted on the live feed's articles: of three thousand seven hundred and forty-seven, the nightly model kept two hundred and sixteen. jev, at the cut-off that catches ninety-five in a hundred on our validation articles, kept eight hundred and ninety, and caught two hundred and thirteen of those two hundred and sixteen. So the next stage would run on four times as many articles. At a stricter cut-off it keeps three hundred and eighty-three and catches only eighty-four in a hundred. The cut-off is a number somebody chooses, and nobody reviews it.

Extraction and retrieval are the same model and the same code in both, because jev cannot write.

Resolution, which record is it, on two thousand two hundred and eighty-one stored entities not used to tune jev: the stored answers say sixty-one in a hundred the same record, thirty-one new, eight unsure. jev, with code for the entities that have no candidate, says sixty-one, thirty-eight, one.

Time: the reviewer takes a median thirty-three seconds a card, fifty-five minutes for forty-eight cards. jev answers forty-one to seventy decisions a second.

The last row is the only place a person's decisions exist: the same forty-eight cards. The reviewer amended thirty-nine, made six new records and judged three not worth a record. Without a person: eight amended, eighteen with no change, twelve new records, one not worth a record, one kept for review, and eight dropped unseen. Had the stored proposals been accepted unreviewed, they would have matched the reviewer on forty-three of the forty-eight, so most of the gap is jev, not the missing person. Forty-one of those forty-eight proposals came from the archive's internal model. Careful: these cards were chosen to include cases where the models disagreed, so this overstates how often they differ.

Two more numbers from the whole data set. In our first audit, of the changes both models agreed on, the reviewer rejected twenty-two in a hundred, all descriptions. And ten of the two thousand seven hundred and seventy-eight matched records would have become duplicate new records.

A note on names: the relevance answers are the nightly Gemini model's. Of the stored resolutions, nine in ten came from the archive's internal model and the rest from Gemini, so "now" means the pipeline's stored answers, not one model. This is a pilot on stored data, with one reviewer who is not a cataloguer. It measures agreement, not accuracy.""",
    21: """Only if asked. The short version of the same comparison. Relevance: of three thousand seven hundred and forty-seven live articles, the nightly model kept 5.8 per cent; jev, at a cut-off that keeps nearly all of those, kept 23.75 per cent, four times the work for the next stage. Resolution, on two thousand two hundred and eighty-one stored entities: the stored answers say sixty-one in a hundred the same record, thirty-one new, eight unsure; jev says sixty-one, thirty-eight, one. Nine in ten of those stored answers came from the archive's internal model, not Gemini.

Decision time: thirty-three seconds a card for the reviewer, milliseconds for jev. On the forty-eight pilot cards, jev's decisions match the reviewer on sixteen; the stored proposals, accepted unreviewed, on forty-three. One reviewer, who is not a cataloguer, on cards chosen partly where the models disagreed: agreement, not accuracy.""",
    # the archive was read by the internal model, not gemini-3.8-flash
    22: """Only if asked. These are from 28 September 2026. The archive is the three outlets' own archives since 2015, read once by an internal model; the live feed has run through Gemini since 3 August. Nothing has been decided by a reviewer in a formal evaluation. About 6,000 matches had nothing to add, and 22,331 entities have no TTE record at all.""",
}


def words(text: str) -> int:
    return len(text.split())


def main() -> int:
    prs = Presentation(str(DECK))
    before = {i: (s.notes_slide.notes_text_frame.text if s.has_notes_slide else "") for i, s in enumerate(prs.slides, 1)}
    assert len(before) == 29, f"expected the 29-slide deck, found {len(before)}"
    for n, text in NOTES.items():
        prs.slides[n - 1].notes_slide.notes_text_frame.text = text
    prs.save(str(OUT_DECK))
    after = {i: (s.notes_slide.notes_text_frame.text if s.has_notes_slide else "") for i, s in enumerate(Presentation(str(OUT_DECK)).slides, 1)}
    untouched = [i for i in after if i not in NOTES]
    assert all(after[i] == before[i] for i in untouched), "a slide's notes changed that should not have"

    d = docx.Document()
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading("Designing for the Reviewer: talking points", 0)
    d.add_paragraph(f"IFLA AI symposium: about 15 minutes of talking, then 5 for questions. The clock is a running time at "
                    f"{WPM} words a minute. Slides after the thank-you slide are backup, for questions only.")
    clock, talk_words = 0.0, 0
    for i in range(1, len(after) + 1):
        text = after[i].strip()
        if i <= TALK_SLIDES:
            stamp = f"[{int(clock // 60)}:{int(clock % 60):02d}]"
            clock += words(text) / WPM * 60 + PAUSES.get(i, 0)
            talk_words += words(text)
        else:
            stamp = "[backup]"
        d.add_heading(f"Slide {i}: {TITLES[i]}  {stamp}", level=2)
        for chunk in text.split("\n\n"):
            if chunk.strip():
                d.add_paragraph(chunk.strip())
    d.add_paragraph(f"Talk: {talk_words} words, about {int(clock // 60)} min {int(clock % 60):02d} s at {WPM} words a minute.")
    d.save(str(OUT_DOC))
    print(f"Wrote {OUT_DECK.relative_to(ROOT)} (notes changed on slides {', '.join(map(str, NOTES))}; slides themselves untouched)")
    print(f"Wrote {OUT_DOC.relative_to(ROOT)}")
    print(f"Talk (slides 1-{TALK_SLIDES}): {words(' '.join(before[i] for i in range(1, TALK_SLIDES + 1)))} words before, "
          f"{talk_words} now; about {int(clock // 60)} min {int(clock % 60):02d} s at {WPM} words a minute")
    return 0


if __name__ == "__main__":
    sys.exit(main())
